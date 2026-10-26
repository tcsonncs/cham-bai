# -*- coding: utf-8 -*-
"""App chấm bài Toán — giao diện một cột cho điện thoại.

Chạy:  streamlit run app.py
"""
import io
import json
from datetime import datetime
from pathlib import Path

import streamlit as st

import cham_bai
import doc_bai
import luu_tru
from cau_hinh import chuan_bi_thu_muc, doc_khoa
from xuat_file import tao_file_excel

st.set_page_config(page_title="Chấm bài Toán", page_icon="📝", layout="centered")

st.markdown(
    """
    <style>
      .block-container {padding-top: 1rem; padding-bottom: 3rem; max-width: 720px;}
      .stButton > button {min-height: 44px; font-size: 16px;}
      textarea, input {font-size: 16px !important;}
      div[data-testid="stMetricValue"] {font-size: 2rem;}
    </style>
    """,
    unsafe_allow_html=True,
)

chuan_bi_thu_muc()


# ---------------- tiện ích ----------------

def lay_dot_hien_tai():
    dot_id = st.session_state.get("dot_id")
    return luu_tru.lay_dot(dot_id) if dot_id else None


def doc_file_word(tep) -> str:
    """Bóc chữ từ .docx (đoạn văn và bảng). Các định dạng khác đọc như văn bản."""
    ten = (tep.name or "").lower()
    if ten.endswith(".docx"):
        import docx

        tai_lieu = docx.Document(io.BytesIO(tep.getvalue()))
        phan = [p.text for p in tai_lieu.paragraphs if p.text.strip()]
        for bang in tai_lieu.tables:
            for dong in bang.rows:
                phan.append(" | ".join(o.text.strip() for o in dong.cells))
        return "\n".join(phan)
    return tep.getvalue().decode("utf-8", errors="ignore")


def hien_thi_diem(ket_qua: dict, khoa: str):
    for cau in ket_qua.get("cau", []):
        co = " ⚠️" if cau.get("nghi_van") else ""
        tieu_de = (f"{cau.get('ma')} — {cau.get('diem_chot')}/"
                   f"{cau.get('diem_toi_da')}{co}")
        with st.expander(tieu_de):
            if cau.get("nghi_van") and cau.get("ly_do_nghi_van"):
                st.warning(cau["ly_do_nghi_van"])
            if cau.get("chon"):
                st.write(f"Đã chọn: **{', '.join(cau['chon'])}** — đáp án "
                         f"**{cau.get('dap_an')}**")
            if cau.get("chi_tiet"):
                for ct in cau["chi_tiet"]:
                    dau = "✅" if ct.get("dung") else "❌"
                    st.write(f"{dau} {ct['y']}) chọn {ct.get('chon') or '(trống)'}"
                             f" — đáp án {ct.get('dap_an')}")
            for buoc in cau.get("buoc") or []:
                st.write(f"- {buoc.get('mo_ta')}: **{buoc.get('diem')}**/"
                         f"{buoc.get('diem_toi_da')} {buoc.get('nhan_xet', '')}")
            if cau.get("nhan_xet"):
                st.info(cau["nhan_xet"])
            buoc = 0.1 if cau.get("loai") == "dung_sai" else 0.25
            cau["diem_chot"] = st.number_input(
                "Điểm chốt cho câu này", min_value=0.0,
                max_value=float(cau.get("diem_toi_da") or 0),
                value=float(cau.get("diem_chot") or 0), step=buoc,
                key=f"{khoa}_{cau.get('ma')}",
            )


# ---------------- thanh điều hướng ----------------

dot = lay_dot_hien_tai()
if not doc_khoa():
    st.error("Chưa có GEMINI_API_KEY. Tạo tệp .env với dòng "
             "GEMINI_API_KEY=... rồi khởi động lại app.")
    st.stop()

st.title("Chấm bài Toán")
if dot:
    st.caption(f"Đợt: **{dot['ten']}** — {dot.get('lop') or ''}")

trang = st.radio(
    "Chọn việc", ["Đợt & đề", "Học sinh", "Ảnh & đọc", "Chấm & duyệt", "Xuất"],
    horizontal=True, label_visibility="collapsed",
)

# ---------------- Đợt & đề ----------------

if trang == "Đợt & đề":
    ds_dot = luu_tru.danh_sach_dot()
    if ds_dot:
        nhan = {f"{d['ten']} (id {d['id']})": d["id"] for d in ds_dot}
        dang_chon = next((k for k, v in nhan.items()
                          if v == st.session_state.get("dot_id")), None)
        chon = st.selectbox("Đợt chấm", list(nhan), index=(
            list(nhan).index(dang_chon) if dang_chon in nhan else 0))
        if st.button("Mở đợt này"):
            st.session_state["dot_id"] = nhan[chon]
            st.rerun()

    with st.expander("Tạo đợt mới"):
        ten_moi = st.text_input("Tên đợt", placeholder="VD: Ôn tập GK1 Toán 9 số 4")
        lop_moi = st.text_input("Lớp", placeholder="VD: 9A")
        cap_moi = st.selectbox("Cấp học", ["THPT", "Đại học"])
        if st.button("Tạo đợt"):
            if ten_moi.strip():
                st.session_state["dot_id"] = luu_tru.tao_dot(
                    ten_moi.strip(), lop_moi.strip(), cap_moi)
                st.rerun()
            else:
                st.warning("Cần nhập tên đợt.")

    dot = lay_dot_hien_tai()
    if not dot:
        st.info("Chưa có đợt nào. Tạo đợt mới để bắt đầu.")
        st.stop()

    st.subheader("Đề và biểu điểm")
    tep = st.file_uploader("Tải file Word đề + biểu điểm", type=["docx", "txt", "md"])
    noi_dung = st.text_area(
        "Nội dung đề + đáp án + biểu điểm (sửa được)",
        value=dot.get("de_bai") or "", height=260,
        placeholder="Dán đề và biểu điểm vào đây, hoặc tải file Word ở trên.",
    )
    if tep and st.button("Bóc chữ từ file vừa tải"):
        st.session_state["de_tam"] = doc_file_word(tep)
        st.rerun()
    if st.session_state.get("de_tam"):
        noi_dung = st.text_area("Nội dung bóc từ file (sửa rồi lưu)",
                                value=st.session_state["de_tam"], height=260)

    if st.button("Lưu đề"):
        luu_tru.cap_nhat_dot(dot["id"], de_bai=noi_dung)
        st.success("Đã lưu đề.")

    st.subheader("Rubric")
    if dot.get("rubric"):
        try:
            rubric = json.loads(dot["rubric"])
            tong = rubric.get("tong_diem")
            st.write(f"Đã có rubric: **{len(rubric.get('cau', []))} câu**, "
                     f"tổng {tong} điểm.")
            for cau in rubric.get("cau", []):
                st.write(f"- {cau.get('ma')} ({cau.get('loai')}) — "
                         f"{cau.get('diem')} điểm")
            with st.expander("Xem/sửa JSON rubric"):
                sua = st.text_area("Rubric JSON", value=dot["rubric"], height=320)
                if st.button("Lưu rubric đã sửa"):
                    json.loads(sua)
                    luu_tru.cap_nhat_dot(dot["id"], rubric=sua)
                    st.success("Đã lưu.")
        except (json.JSONDecodeError, TypeError):
            st.error("Rubric đã lưu không phải JSON hợp lệ.")
    else:
        st.caption("Chưa có rubric. Cần đề + biểu điểm ở trên rồi bấm nút dưới.")
        if st.button("Nhờ AI lập rubric"):
            if not noi_dung.strip():
                st.warning("Chưa có nội dung đề.")
            else:
                with st.spinner("Đang lập rubric..."):
                    try:
                        rubric = cham_bai.tao_rubric(noi_dung)
                        st.session_state["rubric_tam"] = json.dumps(
                            rubric, ensure_ascii=False, indent=2)
                    except Exception as exc:  # noqa: BLE001
                        st.error(f"Lỗi: {exc}")

    if st.session_state.get("rubric_tam"):
        st.subheader("Kiểm tra rubric trước khi dùng")
        sua = st.text_area("Rubric do AI lập (sửa được)",
                           value=st.session_state["rubric_tam"], height=360)
        if st.button("Dùng rubric này"):
            try:
                json.loads(sua)
                luu_tru.cap_nhat_dot(dot["id"], rubric=sua,
                                     de_bai=noi_dung)
                st.session_state.pop("rubric_tam", None)
                st.success("Đã lưu rubric. Sang mục Học sinh để nhập danh sách.")
            except json.JSONDecodeError as exc:
                st.error(f"JSON lỗi: {exc}")

# ---------------- Học sinh ----------------

elif trang == "Học sinh":
    dot = lay_dot_hien_tai()
    if not dot:
        st.warning("Chọn hoặc tạo đợt trước.")
        st.stop()

    ds = luu_tru.danh_sach_hoc_sinh(dot["id"])
    st.write(f"Đợt này có **{len(ds)}** học sinh.")
    for hs in ds:
        cot1, cot2 = st.columns([4, 1])
        cot1.write(hs["ho_ten"])
        if cot2.button("Xoá", key=f"xoa_hs_{hs['id']}"):
            luu_tru.xoa_hoc_sinh(hs["id"])
            st.rerun()

    st.subheader("Thêm học sinh")
    danh_sach = st.text_area("Mỗi dòng một họ tên", height=160,
                             placeholder="Nguyễn Văn A\nTrần Thị B")
    if st.button("Thêm vào đợt"):
        dem = 0
        for dong in danh_sach.splitlines():
            if dong.strip():
                luu_tru.them_hoc_sinh(dot["id"], dong.strip())
                dem += 1
        if dem:
            st.success(f"Đã thêm {dem} học sinh.")
            st.rerun()

# ---------------- Ảnh & đọc ----------------

elif trang == "Ảnh & đọc":
    dot = lay_dot_hien_tai()
    if not dot:
        st.warning("Chọn hoặc tạo đợt trước.")
        st.stop()
    ds = luu_tru.danh_sach_hoc_sinh(dot["id"])
    if not ds:
        st.warning("Chưa có học sinh. Sang mục Học sinh để nhập.")
        st.stop()

    ten_hs = st.selectbox("Chọn học sinh", [h["ho_ten"] for h in ds])
    hs = next(h for h in ds if h["ho_ten"] == ten_hs)
    bai = luu_tru.lay_bai(dot["id"], hs["id"])

    st.subheader("Ảnh bài làm")
    tai_len = st.file_uploader("Chọn ảnh từ điện thoại (chọn được nhiều ảnh)",
                               type=["jpg", "jpeg", "png", "webp"],
                               accept_multiple_files=True)
    chup = st.camera_input("Hoặc chụp ảnh ngay")
    if st.button("Lưu ảnh vào bài"):
        ds_anh = []
        for tep in tai_len or []:
            duong_dan = luu_tru.luu_file_anh(dot["id"], hs["id"], tep.name,
                                             tep.getvalue())
            ds_anh.append({"ten": tep.name, "duong_dan": duong_dan})
        if chup is not None:
            duong_dan = luu_tru.luu_file_anh(dot["id"], hs["id"], "camera.jpg",
                                             chup.getvalue())
            ds_anh.append({"ten": "camera.jpg", "duong_dan": duong_dan})
        if ds_anh:
            luu_tru.luu_anh_bai(dot["id"], hs["id"], ds_anh)
            st.success(f"Đã lưu {len(ds_anh)} ảnh.")
            st.rerun()

    bai = luu_tru.lay_bai(dot["id"], hs["id"])
    if not bai:
        st.info("Chưa có ảnh cho học sinh này.")
        st.stop()
    ds_anh = luu_tru.lay_anh(bai)
    st.caption(f"Bài này có {len(ds_anh)} ảnh.")
    for duong_dan in ds_anh:
        st.image(duong_dan, use_container_width=True)
    if not ds_anh and bai.get("text_doc"):
        st.info("Ảnh của bài này không còn trên máy chủ (máy chủ mây xoá ổ đĩa "
                "khi khởi động lại). Nội dung đã bóc vẫn còn, chấm được bình "
                "thường; muốn đọc lại thì tải ảnh lên lại.")

    if ds_anh and st.button("Đọc bài bằng AI"):
        thanh = st.progress(0.0, text="Đang đọc...")

        def tien_do(i, tong, ten):
            thanh.progress(i / tong, text=f"Đang đọc ảnh {i}/{tong}")

        du_lieu = [(Path(p).name, Path(p).read_bytes()) for p in ds_anh]
        try:
            ket_qua = doc_bai.doc_ca_bai(du_lieu, tien_do=tien_do)
            luu_tru.cap_nhat_bai(
                bai["id"], text_doc=ket_qua,
                trac_nghiem={
                    "trac_nghiem": ket_qua.get("trac_nghiem"),
                    "dung_sai": ket_qua.get("dung_sai"),
                    "tra_loi_ngan": ket_qua.get("tra_loi_ngan"),
                },
                trang_thai="da_doc")
            thanh.empty()
            st.success("Đọc xong. Kiểm tra nội dung bên dưới rồi xác nhận.")
            st.rerun()
        except Exception as exc:  # noqa: BLE001
            thanh.empty()
            st.error(f"Lỗi khi đọc: {exc}")

    if bai.get("text_doc"):
        bai = luu_tru.doc_bai_day_du(bai)
        st.subheader("Kiểm tra nội dung bóc được")
        st.caption("So với ảnh gốc phía trên. Sửa nếu AI đọc sai.")
        goc = bai.get("text_chot") or doc_bai.gop_text(bai["text_doc"])
        sua = st.text_area("Nội dung bài làm", value=goc, height=320)

        marks = bai.get("trac_nghiem") or {}
        st.write("**Trắc nghiệm đọc được** (sửa được nếu máy đọc sai)")
        import pandas as pd

        rows = marks.get("trac_nghiem") or []
        if rows:
            bang = pd.DataFrame([
                {"Câu": r.get("cau"),
                 "Đã chọn": ", ".join(r.get("chon") or []),
                 "Chắc chắn": bool(r.get("chac", True))}
                for r in rows
            ])
            sua_bang = st.data_editor(bang, hide_index=True,
                                      key=f"tn_{bai['id']}")
            rows_moi = []
            for _, r in sua_bang.iterrows():
                chon = [x.strip().upper()
                        for x in str(r["Đã chọn"]).replace(";", ",").split(",")
                        if x.strip()]
                rows_moi.append({"cau": int(r["Câu"]), "chon": chon,
                                 "chac": bool(r["Chắc chắn"])})
            marks["trac_nghiem"] = rows_moi

        ds_ds = marks.get("dung_sai") or []
        if ds_ds:
            st.write("**Câu đúng/sai**")
            for item in ds_ds:
                st.write(f"Câu {item.get('cau', '')} ý {item.get('y')}: "
                         f"{item.get('chon')}")
        for item in marks.get("tra_loi_ngan") or []:
            if isinstance(item, dict):
                st.write(f"Trả lời ngắn câu {item.get('cau')}: "
                         f"**{item.get('tra_loi')}**")

        if st.button("Xác nhận nội dung"):
            luu_tru.cap_nhat_bai(bai["id"], text_chot=sua, trac_nghiem=marks,
                                 trang_thai="da_xac_nhan")
            st.success("Đã xác nhận. Sang mục Chấm & duyệt.")

# ---------------- Chấm & duyệt ----------------

elif trang == "Chấm & duyệt":
    dot = lay_dot_hien_tai()
    if not dot:
        st.warning("Chọn hoặc tạo đợt trước.")
        st.stop()
    if not dot.get("rubric"):
        st.warning("Chưa có rubric. Vào mục Đợt & đề để lập trước.")
        st.stop()
    dst = [h for h in luu_tru.danh_sach_hoc_sinh(dot["id"])
           if luu_tru.lay_bai(dot["id"], h["id"])]
    if not dst:
        st.warning("Chưa có bài nào được nạp ảnh.")
        st.stop()

    ten_hs = st.selectbox("Chọn học sinh", [h["ho_ten"] for h in dst])
    hs = next(h for h in dst if h["ho_ten"] == ten_hs)
    bai = luu_tru.doc_bai_day_du(luu_tru.lay_bai(dot["id"], hs["id"]))
    rubric = json.loads(dot["rubric"])

    if not bai.get("text_doc"):
        st.warning("Bài này chưa đọc ảnh. Vào mục Ảnh & đọc trước.")
        st.stop()

    if st.button("Chấm bài"):
        du_lieu_doc = dict(bai.get("text_doc") or {})
        du_lieu_doc["text_chot"] = bai.get("text_chot")
        if bai.get("trac_nghiem"):
            du_lieu_doc.update({
                "trac_nghiem": bai["trac_nghiem"].get("trac_nghiem"),
                "dung_sai": bai["trac_nghiem"].get("dung_sai"),
                "tra_loi_ngan": bai["trac_nghiem"].get("tra_loi_ngan"),
            })
        with st.spinner("Đang chấm..."):
            try:
                ket_qua = cham_bai.cham_ca_bai(rubric, du_lieu_doc)
                luu_tru.cap_nhat_bai(
                    bai["id"], ket_qua=ket_qua,
                    tong_goi_y=ket_qua["tong_goi_y"],
                    tong_chot=ket_qua["tong_chot"],
                    trang_thai="da_cham",
                    ngay_cham=datetime.now().isoformat(timespec="seconds"))
                st.rerun()
            except Exception as exc:  # noqa: BLE001
                st.error(f"Lỗi khi chấm: {exc}")

    ket_qua = bai.get("ket_qua")
    if not ket_qua:
        st.info("Bấm Chấm bài để bắt đầu.")
        st.stop()

    nghi_van = [c for c in ket_qua.get("cau", []) if c.get("nghi_van")]
    cot1, cot2 = st.columns(2)
    cot1.metric("Điểm máy gợi ý", ket_qua.get("tong_goi_y"))
    cot2.metric("Điểm chốt", ket_qua.get("tong_chot"))
    st.write(f"Có **{len(nghi_van)}** chỗ cần xem lại.")

    if ket_qua.get("nhan_xet_chung"):
        st.info(ket_qua["nhan_xet_chung"])

    hien_thi_diem(ket_qua, f"cham_{bai['id']}_{bai.get('ngay_cham') or ''}")

    if st.button("Lưu & duyệt bài này"):
        ket_qua["tong_chot"] = round(
            sum(float(c.get("diem_chot") or 0) for c in ket_qua.get("cau", [])), 2)
        luu_tru.cap_nhat_bai(bai["id"], ket_qua=ket_qua,
                             tong_chot=ket_qua["tong_chot"],
                             trang_thai="da_duyet")
        st.success(f"Đã duyệt. Điểm chốt: {ket_qua['tong_chot']}")

    with st.expander("Xem bài làm đã bóc"):
        st.text(bai.get("text_chot") or doc_bai.gop_text(bai.get("text_doc") or {}))

# ---------------- Xuất ----------------

elif trang == "Xuất":
    dot = lay_dot_hien_tai()
    if not dot:
        st.warning("Chọn hoặc tạo đợt trước.")
        st.stop()
    ds = luu_tru.danh_sach_hoc_sinh(dot["id"])
    da_cham = [h for h in ds
               if (luu_tru.lay_bai(dot["id"], h["id"]) or {}).get("ket_qua")]
    st.write(f"Đợt **{dot['ten']}**: {len(da_cham)}/{len(ds)} bài đã chấm.")
    if da_cham:
        du_lieu = tao_file_excel(dot, ds)
        st.download_button("Tải bảng điểm Excel", data=du_lieu,
                           file_name=f"BangDiem_{dot['ten']}.xlsx",
                           mime="application/vnd.openxmlformats-officedocument"
                                ".spreadsheetml.sheet")
    st.caption("Xuất PDF từng học sinh và tin nhắn Zalo sẽ làm ở bước sau.")
